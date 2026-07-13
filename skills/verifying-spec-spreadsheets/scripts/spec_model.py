"""
spec_model.py — Parse a homebuilder "Forward Sale Spec List" workbook into a
flat, structured model and identify which communities are blank/incomplete.

The source workbook uses a hierarchical, merged-cell layout:

    Header row (per community):  Division | Property Name | Type | Lat/Long | Website | Delivery Date
    Detail rows (per floor plan):  Base Price | Floor Plan | Unit Type (beds/baths) |
                                   Lot Width | Product | SF | Stories | Garages
    Community-level (first detail row): Amenities | HOA

Merged cells make in-place row insertion fragile, so this module reads the
hierarchical sheet into plain dataclasses. Downstream tooling rebuilds a clean,
un-merged workbook from the model — see build_output.py.

No external dependency beyond openpyxl (zero-dependency-friendly).
"""
from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field, asdict
from typing import Any, Optional

import openpyxl

# Column map for the BTR Forward Sale Spec layout (1-based indexes).
COL = {
    "division": 2,      # B
    "property": 3,      # C
    "type": 4,          # D
    "latlong": 5,       # E
    "website": 6,       # F
    "delivery": 7,      # G
    "price": 8,         # H
    "plan": 9,          # I
    "unit_type": 10,    # J  (beds/baths, e.g. "3 Bed / 2.5 Bath")
    "lot_width": 11,    # K
    "product": 12,      # L  (Townhome / Single Family / Paired / Duplex)
    "sf": 13,           # M
    "stories": 14,      # N
    "garages": 15,      # O
    "amenities": 16,    # P
    "hoa": 17,          # Q
}

HEADER_ROW = 4          # the labels row
FIRST_DATA_ROW = 5


@dataclass
class FloorPlan:
    base_price: Any = None
    plan: Optional[str] = None
    unit_type: Optional[str] = None      # beds/baths
    lot_width: Optional[str] = None
    product: Optional[str] = None
    sf: Any = None
    stories: Any = None
    garages: Any = None
    source: Optional[str] = None         # provenance for enriched rows
    confidence: Optional[str] = None     # high / medium / low / None=as-given

    def is_empty(self) -> bool:
        return not any([self.base_price, self.plan, self.unit_type, self.sf])


@dataclass
class Community:
    src_row: int
    division: Optional[str] = None
    property_name: Optional[str] = None
    type: Optional[str] = None
    latlong: Optional[str] = None
    website: Optional[str] = None
    delivery: Any = None
    amenities: Optional[str] = None
    hoa: Optional[str] = None
    floor_plans: list = field(default_factory=list)
    # Enrichment / verification metadata (populated by tooling, not the source)
    status: str = "complete"             # complete | blank | enriched | needs_research
    source: Optional[str] = None
    notes: Optional[str] = None

    def missing_fields(self) -> list:
        m = []
        if not self.latlong:
            m.append("latlong")
        if not self.website:
            m.append("website")
        if not self.floor_plans or all(fp.is_empty() for fp in self.floor_plans):
            m.append("floor_plans")
        if not self.delivery:
            m.append("delivery")
        return m

    def is_blank(self) -> bool:
        # "Blank" = the three substantive fields are all missing.
        m = set(self.missing_fields())
        return {"latlong", "website", "floor_plans"}.issubset(m)


def _clean(v):
    if isinstance(v, str):
        return v.strip() or None
    return v


def load_communities(path: str) -> list:
    """Read the workbook into a list[Community], carrying division forward."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    def cell(r, key):
        return _clean(ws.cell(row=r, column=COL[key]).value)

    # A community block starts on any row with a Property Name (col C).
    starts = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if cell(r, "property")]
    starts.append(ws.max_row + 1)  # sentinel

    communities: list = []
    last_division = None
    for i in range(len(starts) - 1):
        r0, r1 = starts[i], starts[i + 1]
        division = cell(r0, "division") or last_division
        last_division = division
        c = Community(
            src_row=r0,
            division=division,
            property_name=cell(r0, "property"),
            type=cell(r0, "type"),
            latlong=cell(r0, "latlong"),
            website=cell(r0, "website"),
            delivery=cell(r0, "delivery"),
            amenities=cell(r0, "amenities"),
            hoa=cell(r0, "hoa"),
        )
        for r in range(r0, r1):
            fp = FloorPlan(
                base_price=cell(r, "price"),
                plan=cell(r, "plan"),
                unit_type=cell(r, "unit_type"),
                lot_width=cell(r, "lot_width"),
                product=cell(r, "product"),
                sf=cell(r, "sf"),
                stories=cell(r, "stories"),
                garages=cell(r, "garages"),
            )
            if not fp.is_empty():
                c.floor_plans.append(fp)
            # amenities/hoa may live on a non-first detail row
            if not c.amenities:
                c.amenities = cell(r, "amenities")
            if not c.hoa:
                c.hoa = cell(r, "hoa")
        c.status = "blank" if c.is_blank() else "complete"
        communities.append(c)
    return communities


def summarize(communities: list) -> dict:
    blanks = [c for c in communities if c.is_blank()]
    by_div: dict = {}
    for c in communities:
        by_div.setdefault(c.division, {"total": 0, "blank": 0})
        by_div[c.division]["total"] += 1
        if c.is_blank():
            by_div[c.division]["blank"] += 1
    return {
        "total": len(communities),
        "blank": len(blanks),
        "complete": len(communities) - len(blanks),
        "by_division": by_div,
        "blank_list": [
            {"row": c.src_row, "division": c.division, "property": c.property_name,
             "missing": c.missing_fields()}
            for c in blanks
        ],
    }


if __name__ == "__main__":
    import json
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else None
    if not path:
        print("usage: python spec_model.py <workbook.xlsx>")
        raise SystemExit(2)
    comms = load_communities(path)
    print(json.dumps(summarize(comms), indent=2, default=str))
